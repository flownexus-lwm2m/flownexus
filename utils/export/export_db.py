import paramiko
import os
import shutil
import getpass
import os
import sys
import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
import django
from django.conf import settings
from pathlib import Path
import matplotlib.pyplot as plt
import numpy as np


# Export Directory
EXPORT_PUMP_EVTS = Path('./export_pump_evts')

endpoints_include = ["358299840022335"]
days_last_include = 10


# Django Environment Setup
script_dir = os.path.dirname(os.path.abspath(__file__))
project_dir = os.path.abspath(os.path.join(script_dir, '../../server/django/'))
sys.path.insert(0, project_dir)
os.environ.setdefault("DJANGO_SETTINGS_MODULE", "server.settings")
# Override DATABASES setting dynamically before setup
db_path = Path('./db.sqlite3').resolve()

if not settings.configured:
    settings.configure(
        DATABASES={
            "default": {
                "ENGINE": "django.db.backends.sqlite3",
                "NAME": str(db_path),
            }
        },
        INSTALLED_APPS=[
            "sensordata",
        ],
        # include other minimal required settings if needed
    )
django.setup()
from sensordata.models import Endpoint, Resource

# SSH Configuration
hostname = "85.215.45.193"
port = 22
username = "jonas"
remote_path = "/home/jonas/flownexus/server/django/db.sqlite3"
local_path = "db.sqlite3"


def download_file_via_ssh_default_key(host, p, user, remote_path, local_path):
    ssh = paramiko.SSHClient()
    ssh.set_missing_host_key_policy(paramiko.AutoAddPolicy())

    # Try to use keys from ssh-agent
    agent = paramiko.Agent()
    keys = agent.get_keys()
    if keys:
        # Use the first key from the agent
        private_key = keys[0]
    else:
        print("No SSH keys found in ssh-agent")

    ssh.connect(host, port=p, username=user, pkey=private_key)

    sftp = ssh.open_sftp()
    sftp.get(remote_path, local_path)
    sftp.close()
    ssh.close()


def export_endpoint_resources():
    # Get all endpoints
    endpoints = Endpoint.objects.all()

    rows = []
    for ep in endpoints:
        # Get latest resources for this endpoint
        resources = Resource.objects.filter(endpoint=ep).select_related('resource_type').order_by('-timestamp_created')
        for res in resources:
            rows.append({
                "endpoint": ep.endpoint,
                "registered": ep.registered,
            })

    return pd.DataFrame(rows)


def export_resources():
    resources = Resource.objects.select_related('endpoint', 'resource_type').all()
    print(f"Resource objects: {Resource.objects.count()}")
    rows = []
    for res in resources:
        # Gather event info for this resource
        events = [f"{er.event.event_type} ({er.event.time})" for er in res.eventresource_set.all()]
        events_str = "; ".join(events) if events else None

        rows.append({
            "endpoint": res.endpoint.endpoint,
            "resource_type": str(res.resource_type),
            "timestamp_created": res.timestamp_created.replace(tzinfo=None) if res.timestamp_created else None,
            "int_value": res.int_value,
            "float_value": res.float_value,
            "str_value": res.str_value,
            "binary_value": res.binary_value,
            "value": res.get_value(),
            "events": events_str,
        })
    return pd.DataFrame(rows)


def export_pump_event_resources():
    resources = Resource.objects.select_related(
            'endpoint',
            'resource_type').filter(
        resource_type__object_id=10300
        )

    print(f"Filtered Resource objects: {resources.count()}")
    rows = []
    for res in resources:
        # Gather event info for this resource
        #events = [f"{er.event.event_type} ({er.event.time})" for er in res.eventresource_set.all()]
        #events_str = "; ".join(events) if events else None

        er = res.eventresource_set.first()
        event_ts = er.event.time if er else None

        rows.append({
            "endpoint": res.endpoint.endpoint.split(":")[-1],
            "object_id": res.resource_type.object_id,
            "resource_id": res.resource_type.resource_id,
            "timestamp_created": res.timestamp_created.replace(tzinfo=None) if res.timestamp_created else None,
            "int_value": res.int_value,
            "binary_value": res.binary_value,
            "event_ts": event_ts.replace(tzinfo=None) if event_ts else None,
        })
    return pd.DataFrame(rows)


def get_endpoints():
    endpoints = Endpoint.objects.all()
    rows = []
    for ep in endpoints:
        rows.append({
            "endpoint": ep.endpoint.split(":")[-1],
            "registered": ep.registered,
        })
    print(f"Endpoint objects {Endpoint.objects.count()}")

    return pd.DataFrame(rows)


def add_autofilter(filename, sheet_name):
    wb = load_workbook(filename)
    ws = wb[sheet_name]
    max_col = ws.max_column
    max_row = ws.max_row
    ws.auto_filter.ref = f"A1:{get_column_letter(max_col)}{max_row}"
    wb.save(filename)


def export_db_xslx():
    with pd.ExcelWriter('export.xlsx', engine='openpyxl') as writer:

        df = export_resources()
        df.to_excel(writer, sheet_name='Resources', index=False)
    add_autofilter('export.xlsx', 'Resources')
    print("Export to xlsx file done.")


def export_graph(df, export_path, ts, ep):
    """
    Plots sensor data with different sampling rates and saves the figure.

    Args:
        df (pd.DataFrame): DataFrame with index as time or sample number and
        columns 'pm_press_avg', 'pm_press_rsd', 'pm_press_raw'.
        export_path (str): Path to save the generated plot.
    """

    # Generate individual time-axis according to sampling rate.
    t_avg = np.arange(len(df['pm_press_avg'])) if 'pm_press_avg' in df else None
    t_rsd = np.arange(len(df['pm_press_rsd'])) if 'pm_press_rsd' in df else None
    t_raw = np.arange(len(df['pm_press_raw'])) / 50.0 if 'pm_press_raw' in df else None

    plt.figure(figsize=(12, 6))

    # Plot pm_press_avg
    if 'pm_press_avg' in df:
        t_avg = np.arange(len(df['pm_press_avg']))
        plt.plot(
            t_avg, df['pm_press_avg'],
            label='pm_press_avg',
            linestyle='-',
            marker='o',
            markersize=5
        )

    # Plot pm_press_rsd
    if 'pm_press_rsd' in df:
        t_rsd = np.arange(len(df['pm_press_rsd']))
        plt.plot(
            t_rsd, df['pm_press_rsd'],
            label='pm_press_rsd',
            linestyle='-',
            marker='o',
            markersize=5
        )

    # Plot pm_press_raw
    if 'pm_press_raw' in df:
        t_raw = np.arange(len(df['pm_press_raw'])) / 50.0
        plt.plot(
            t_raw, df['pm_press_raw'],
            label='pm_press_raw',
            linestyle='-',
            marker='o',
            markersize=3
        )

    plt.xlabel('Time [s]')
    plt.ylabel('Pressure [kPa]   |   RSD Signal')
    plt.legend()
    plt.title(f'Endpoint {ep} - {ts}')
    plt.tight_layout()
    plt.savefig(export_path)
    plt.close()


def export_db():
    if os.path.exists(EXPORT_PUMP_EVTS):
        shutil.rmtree(EXPORT_PUMP_EVTS)
    os.makedirs(EXPORT_PUMP_EVTS, exist_ok=True)

    #res = export_resources()
    res = export_pump_event_resources()

    # Filter resources based on the last days to include
    threshold_date = pd.Timestamp.now() - pd.Timedelta(days=days_last_include)
    filtered_timestamps = res['event_ts'].dropna()
    filtered_timestamps = filtered_timestamps[filtered_timestamps >= threshold_date]

    event_timestamps = filtered_timestamps.unique()
    print(f"Unique events: {len(event_timestamps)} ")

    # Create a folder for each endpoint
    ep = get_endpoints()
    for endpoint in ep['endpoint']:
        # Select individual endpoints to make the export faster
        if endpoint not in endpoints_include:
            continue

        endpoint_dir = os.path.join(EXPORT_PUMP_EVTS, endpoint)
        os.makedirs(endpoint_dir, exist_ok=True)

        for event_ts in event_timestamps:
            event_rows = res[(res['event_ts'] == event_ts) & (res['endpoint'] == endpoint)]
            if event_rows.empty:
                continue

            # Get all values for resource_id == 2 (pm_press_avg value)
            pm_press_avg = event_rows[event_rows['resource_id'] == 2]['int_value'].tolist()
            pm_press_rsd = event_rows[event_rows['resource_id'] == 3]['int_value'].tolist()
            pm_press_raw = event_rows[event_rows['resource_id'] == 4]['binary_value'].tolist()
            pm_press_raw = pm_press_raw[0] if pm_press_raw else None
            if isinstance(pm_press_raw, bytes):
                pm_press_raw = list(pm_press_raw)
                # Convert from pressure to underpressure with an assumed
                # absolute athmospheric pressure of 102 kPa
                pm_press_raw = [102 - x for x in pm_press_raw]
            else:
                pm_press_raw = None

            # Create a DataFrame for the event
            df_event = pd.DataFrame({
                'pm_press_avg': pd.Series(pm_press_avg),
                'pm_press_rsd': pd.Series(pm_press_rsd),
                'pm_press_raw': pd.Series(pm_press_raw),
            })

            # Export to CSV
            filename = f"{event_ts}.csv"
            filepath = os.path.join(endpoint_dir, filename)
            df_event.to_csv(filepath, index=False)

            # Generate Graphs
            export_graph(df_event,
                         os.path.join(endpoint_dir, f"{event_ts}.svg"),
                         event_ts,
                         endpoint)



if __name__ == "__main__":
    try:
        download_file_via_ssh_default_key(hostname, port, username, remote_path, local_path)
        print(f"Successfully downloaded {remote_path} to {local_path}")
    except Exception as e:
        print(f"An error occurred: {e}")

    #export_db_xslx()
    export_db()

